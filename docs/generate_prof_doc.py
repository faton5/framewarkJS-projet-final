from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

OUT = 'docs/Presentation_professeur.docx'
NAVY = '0F172A'
TEAL = '18A796'
MUTED = '66748A'
LIGHT = 'F5F7FA'
LINE = 'D9D9D9'

def shade(cell, color):
    props = cell._tc.get_or_add_tcPr()
    el = OxmlElement('w:shd')
    el.set(qn('w:fill'), color)
    props.append(el)

def borders(table):
    tblPr = table._tbl.tblPr
    border = tblPr.first_child_found_in('w:tblBorders')
    if border is None:
        border = OxmlElement('w:tblBorders')
        tblPr.append(border)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        element = border.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            border.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '4')
        element.set(qn('w:color'), LINE)

def set_cell(cell, text, bold=False, color=NAVY, size=9.5):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = 'Arial'
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    borders(table)
    for i, header in enumerate(headers):
        set_cell(table.rows[0].cells[i], header, True, 'FFFFFF', 9)
        shade(table.rows[0].cells[i], NAVY)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell(cells[i], value)
            if row_index % 2 == 1:
                shade(cells[i], LIGHT)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table

def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p

def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='Normal')
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2.1)
section.right_margin = Cm(2.1)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(NAVY)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.12
for name, size in [('Title', 27), ('Heading 1', 16), ('Heading 2', 11.5)]:
    style = styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor.from_string(NAVY if name != 'Heading 2' else TEAL)
    style.font.size = Pt(size)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(16 if name == 'Heading 1' else 10)
    style.paragraph_format.space_after = Pt(6)

title = doc.add_paragraph(style='Title')
title.add_run('Présentation du projet Campus Matériel')
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(12)
r = subtitle.add_run('Application web de consultation et de réservation de matériel pédagogique')
r.font.name = 'Arial'; r.font.size = Pt(12); r.font.color.rgb = RGBColor.from_string(MUTED)
meta = doc.add_paragraph()
meta.paragraph_format.space_after = Pt(18)
r = meta.add_run('Technologie : React · Vite · React Router · localStorage\nProjet réalisé dans le cadre du projet final Campus Matériel')
r.font.name = 'Arial'; r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string(MUTED)

add_heading(doc, '1. Objectif du projet')
add_body(doc, "Campus Matériel est une application monopage destinée à un établissement qui souhaite proposer son matériel aux étudiants. Elle permet de consulter un catalogue, de réserver un équipement disponible et de gérer les fiches de matériel avec un compte responsable.")
add_body(doc, "La connexion est volontairement simulée avec trois comptes de démonstration. L’objectif est de montrer le fonctionnement frontend, la gestion d’état, la navigation et la persistance locale sans mettre en place de serveur.")

add_heading(doc, '2. Fonctionnalités réalisées')
add_table(doc, ['Partie', 'Fonctionnement'], [
    ('Catalogue', 'Six équipements répartis dans quatre catégories, recherche par nom, filtre par catégorie, indication de disponibilité et fiche détaillée.'),
    ('Réservations', 'Un étudiant peut réserver plusieurs équipements différents. Un équipement ne peut avoir qu’une réservation active. L’annulation libère immédiatement le matériel.'),
    ('Gestion', 'Le responsable peut créer, modifier et supprimer une fiche. La suppression d’un équipement réservé est refusée avec un message explicite.'),
    ('Validation', 'Le formulaire partagé vérifie le nom, la description et la localisation, en conservant les valeurs saisies en cas d’erreur.'),
    ('Persistance', 'La session, le catalogue et les réservations sont restaurés depuis localStorage après rechargement.'),
])

add_heading(doc, '3. Comptes et routes de démonstration')
add_table(doc, ['Compte', 'Rôle', 'Droits principaux'], [
    ('Alice Martin', 'Étudiante', 'Catalogue, détail, réservation, annulation de ses réservations'),
    ('Yanis Bernard', 'Étudiant', 'Catalogue, détail, réservation, annulation de ses réservations'),
    ('Camille Dupont', 'Responsable', 'Catalogue et gestion complète des fiches'),
])
add_body(doc, 'Routes principales : /login, /catalogue, /materiels/:id, /mes-reservations, /gestion, /gestion/nouveau et /gestion/:id/modifier.', 'Routes principales :')

add_heading(doc, '4. Parcours d’une réservation')
add_body(doc, 'Le parcours est entièrement piloté par l’état partagé fourni par AppContext :')
flow = doc.add_paragraph()
flow.alignment = WD_ALIGN_PARAGRAPH.CENTER
flow.paragraph_format.space_before = Pt(5); flow.paragraph_format.space_after = Pt(10)
r = flow.add_run('Clic sur Réserver  →  contrôle de disponibilité  →  ajout  →  sauvegarde  →  mise à jour des vues')
r.bold = True; r.font.color.rgb = RGBColor.from_string(TEAL); r.font.size = Pt(11)
for item in [
    'L’étudiant ouvre la fiche d’un équipement.',
    'L’application cherche une réservation active portant le même identifiant.',
    'Si l’équipement est libre, une réservation est créée avec un identifiant unique, l’identifiant du matériel et celui de l’étudiant.',
    'Le nouvel état est enregistré dans localStorage.',
    'Le catalogue, la fiche et la page « Mes réservations » se mettent à jour sans rechargement manuel.',
]:
    p = doc.add_paragraph(style='List Number')
    p.add_run(item)

add_heading(doc, '5. Organisation technique')
for item in [
    'Pages : écrans de connexion, catalogue, détail, réservations, gestion, accès refusé et 404.',
    'Composants réutilisables : mise en page, cartes équipement, badge de disponibilité et icônes.',
    'État partagé : session, équipements et réservations dans un contexte React.',
    'Protection : les routes privées redirigent vers la connexion et les routes réservées au responsable contrôlent le rôle.',
    'Chargement : l’espace de gestion est chargé à la demande avec lazy et Suspense.',
]:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(item)

add_heading(doc, '6. Limites connues')
add_body(doc, "Cette version est une démonstration frontend. Il n’y a pas de serveur, de véritable authentification ou de synchronisation entre plusieurs navigateurs. Dans une application réelle, les permissions et les règles de réservation devraient aussi être vérifiées côté serveur.")

add_heading(doc, '7. Lancement')
add_body(doc, 'Depuis le dossier du projet :')
for cmd in ['npm install', 'npm run dev']:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run(cmd); r.font.name = 'Courier New'; r.font.size = Pt(10); r.bold = True
footer = doc.add_paragraph()
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.space_before = Pt(18)
r = footer.add_run('Campus Matériel · Documentation de présentation destinée à l’évaluation')
r.font.name = 'Arial'; r.font.size = Pt(8.5); r.font.color.rgb = RGBColor.from_string(MUTED)

doc.save(OUT)
